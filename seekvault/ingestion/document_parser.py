import os
import fitz  # PyMuPDF
import docx  # python-docx
from typing import Dict, Any, List
from seekvault.ingestion.models import ExtractionResult
from seekvault.core.logger import logger

def extract_pdf(file_path: str) -> ExtractionResult:
    """
    Extracts text and metadata from PDF files using PyMuPDF (fitz).
    Does not alter or write to the original PDF file.
    """
    norm_path = os.path.abspath(os.path.realpath(file_path))
    if not os.path.exists(norm_path):
        return ExtractionResult(
            file_path=norm_path,
            file_type="document",
            extracted_text="",
            metadata={},
            extraction_method="pymupdf",
            extraction_status="FAILED",
            error_message="File not found."
        )

    try:
        doc = fitz.open(norm_path)
        page_texts = []
        pages_extracted = 0
        total_pages = len(doc)
        
        pdf_metadata = {
            "page_count": total_pages,
            "title": doc.metadata.get("title", ""),
            "author": doc.metadata.get("author", ""),
            "subject": doc.metadata.get("subject", ""),
            "format": doc.metadata.get("format", "PDF"),
        }
        
        for page_idx in range(total_pages):
            try:
                page = doc.load_page(page_idx)
                text = page.get_text("text")
                if text and text.strip():
                    page_texts.append(f"--- Page {page_idx + 1} ---\n{text.strip()}")
                    pages_extracted += 1
            except Exception as pe:
                logger.warning(f"PyMuPDF: Failed to extract page {page_idx + 1} of {norm_path}: {pe}")

        doc.close()
        
        full_text = "\n\n".join(page_texts)
        pdf_metadata["char_count"] = len(full_text)
        pdf_metadata["pages_with_text"] = pages_extracted

        if not full_text.strip():
            return ExtractionResult(
                file_path=norm_path,
                file_type="document",
                extracted_text="",
                metadata=pdf_metadata,
                extraction_method="pymupdf",
                extraction_status="EMPTY",
                error_message="PDF contains no selectable text (scanned PDF or blank pages)."
            )
            
        status = "SUCCESS" if pages_extracted == total_pages else "PARTIAL"
        return ExtractionResult(
            file_path=norm_path,
            file_type="document",
            extracted_text=full_text,
            metadata=pdf_metadata,
            extraction_method="pymupdf",
            extraction_status=status
        )

    except Exception as e:
        logger.error(f"PyMuPDF extraction failed for '{norm_path}': {e}")
        return ExtractionResult(
            file_path=norm_path,
            file_type="document",
            extracted_text="",
            metadata={},
            extraction_method="pymupdf",
            extraction_status="FAILED",
            error_message=str(e)
        )

def extract_docx(file_path: str) -> ExtractionResult:
    """
    Extracts text and metadata from DOCX files using python-docx.
    Does not alter or write to the original DOCX file.
    """
    norm_path = os.path.abspath(os.path.realpath(file_path))
    if not os.path.exists(norm_path):
        return ExtractionResult(
            file_path=norm_path,
            file_type="document",
            extracted_text="",
            metadata={},
            extraction_method="python-docx",
            extraction_status="FAILED",
            error_message="File not found."
        )

    try:
        doc = docx.Document(norm_path)
        paragraphs_text = []
        
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                paragraphs_text.append(t)
                
        # Extract text from tables if any
        table_paragraphs = []
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    table_paragraphs.append(" | ".join(row_text))
                    
        if table_paragraphs:
            paragraphs_text.append("\n--- Tables ---")
            paragraphs_text.extend(table_paragraphs)
            
        full_text = "\n".join(paragraphs_text)
        
        docx_metadata = {
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "word_count": len(full_text.split()),
            "char_count": len(full_text)
        }

        if not full_text.strip():
            return ExtractionResult(
                file_path=norm_path,
                file_type="document",
                extracted_text="",
                metadata=docx_metadata,
                extraction_method="python-docx",
                extraction_status="EMPTY",
                error_message="DOCX document contains no text content."
            )

        return ExtractionResult(
            file_path=norm_path,
            file_type="document",
            extracted_text=full_text,
            metadata=docx_metadata,
            extraction_method="python-docx",
            extraction_status="SUCCESS"
        )

    except Exception as e:
        logger.error(f"python-docx extraction failed for '{norm_path}': {e}")
        return ExtractionResult(
            file_path=norm_path,
            file_type="document",
            extracted_text="",
            metadata={},
            extraction_method="python-docx",
            extraction_status="FAILED",
            error_message=str(e)
        )

def extract_txt(file_path: str) -> ExtractionResult:
    """
    Extracts text from plain text files (.txt, .md) using standard Python file handling.
    Tries utf-8 encoding first, falling back to latin-1 / cp1252.
    """
    norm_path = os.path.abspath(os.path.realpath(file_path))
    if not os.path.exists(norm_path):
        return ExtractionResult(
            file_path=norm_path,
            file_type="document",
            extracted_text="",
            metadata={},
            extraction_method="text_reader",
            extraction_status="FAILED",
            error_message="File not found."
        )

    encodings_to_try = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    text_content = ""
    used_encoding = ""
    success = False

    for enc in encodings_to_try:
        try:
            with open(norm_path, "r", encoding=enc) as f:
                text_content = f.read()
            used_encoding = enc
            success = True
            break
        except UnicodeDecodeError:
            continue
        except Exception as e:
            return ExtractionResult(
                file_path=norm_path,
                file_type="document",
                extracted_text="",
                metadata={},
                extraction_method="text_reader",
                extraction_status="FAILED",
                error_message=str(e)
            )

    if not success:
        return ExtractionResult(
            file_path=norm_path,
            file_type="document",
            extracted_text="",
            metadata={},
            extraction_method="text_reader",
            extraction_status="FAILED",
            error_message="Failed to decode text file with standard encodings."
        )

    lines = text_content.splitlines()
    txt_metadata = {
        "encoding": used_encoding,
        "line_count": len(lines),
        "word_count": len(text_content.split()),
        "char_count": len(text_content)
    }

    if not text_content.strip():
        return ExtractionResult(
            file_path=norm_path,
            file_type="document",
            extracted_text="",
            metadata=txt_metadata,
            extraction_method="text_reader",
            extraction_status="EMPTY",
            error_message="Text file is empty."
        )

    return ExtractionResult(
        file_path=norm_path,
        file_type="document",
        extracted_text=text_content.strip(),
        metadata=txt_metadata,
        extraction_method="text_reader",
        extraction_status="SUCCESS"
    )
