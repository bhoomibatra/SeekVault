import os
import tempfile
import pytest
import fitz  # PyMuPDF
import docx  # python-docx
from PIL import Image, ImageDraw

from seekvault.ingestion.models import ExtractionResult
from seekvault.ingestion.document_parser import extract_pdf, extract_docx, extract_txt
from seekvault.ingestion.image_parser import extract_image
from seekvault.ingestion.extractor import extract_file
from seekvault.ingestion.hasher import compute_sha256

@pytest.fixture
def test_files_dir():
    with tempfile.TemporaryDirectory() as tmpdir:
        # 1. Create PDF
        pdf_path = os.path.join(tmpdir, "sample.pdf")
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "Hello SeekVault PDF Extraction Test Page")
        doc.save(pdf_path)
        doc.close()

        # 2. Create DOCX
        docx_path = os.path.join(tmpdir, "sample.docx")
        doc_w = docx.Document()
        doc_w.add_heading("SeekVault Document Test", level=1)
        doc_w.add_paragraph("This is a paragraph inside the DOCX test file.")
        doc_w.save(docx_path)

        # 3. Create TXT
        txt_path = os.path.join(tmpdir, "sample.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("Line 1: SeekVault plain text extraction test.\nLine 2: Data content.")

        # 4. Create PNG Image
        png_path = os.path.join(tmpdir, "sample.png")
        img = Image.new("RGB", (300, 100), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        draw.text((10, 40), "SEEKVAULT OCR TEST", fill=(0, 0, 0))
        img.save(png_path)

        yield {
            "pdf": pdf_path,
            "docx": docx_path,
            "txt": txt_path,
            "png": png_path
        }

def test_extract_pdf(test_files_dir):
    res = extract_pdf(test_files_dir["pdf"])
    assert isinstance(res, ExtractionResult)
    assert res.file_type == "document"
    assert res.extraction_method == "pymupdf"
    assert res.extraction_status in ["SUCCESS", "PARTIAL"]
    assert "Hello SeekVault PDF" in res.extracted_text
    assert res.metadata["page_count"] == 1

def test_extract_docx(test_files_dir):
    res = extract_docx(test_files_dir["docx"])
    assert isinstance(res, ExtractionResult)
    assert res.file_type == "document"
    assert res.extraction_method == "python-docx"
    assert res.extraction_status == "SUCCESS"
    assert "SeekVault Document Test" in res.extracted_text
    assert res.metadata["paragraph_count"] > 0

def test_extract_txt(test_files_dir):
    res = extract_txt(test_files_dir["txt"])
    assert isinstance(res, ExtractionResult)
    assert res.file_type == "document"
    assert res.extraction_method == "text_reader"
    assert res.extraction_status == "SUCCESS"
    assert "Line 1: SeekVault" in res.extracted_text
    assert res.metadata["line_count"] == 2

def test_extract_image(test_files_dir):
    res = extract_image(test_files_dir["png"])
    assert isinstance(res, ExtractionResult)
    assert res.file_type == "image"
    assert res.metadata["width"] == 300
    assert res.metadata["height"] == 100
    assert res.metadata["format"] == "PNG"
    assert res.extraction_status in ["SUCCESS", "EMPTY"]

def test_unified_extract_file(test_files_dir):
    for key, path in test_files_dir.items():
        res = extract_file(path)
        assert isinstance(res, ExtractionResult)
        assert res.file_path == os.path.abspath(os.path.realpath(path))
        assert res.extraction_status in ["SUCCESS", "PARTIAL", "EMPTY"]
        assert res.extraction_method != "none"

def test_zero_touch_during_extraction(test_files_dir):
    pdf_path = test_files_dir["pdf"]
    initial_hash = compute_sha256(pdf_path)
    
    # Run extraction
    res = extract_file(pdf_path)
    
    post_hash = compute_sha256(pdf_path)
    assert initial_hash == post_hash, "PDF file was modified during extraction!"
