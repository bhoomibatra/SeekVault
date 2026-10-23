import os
from pathlib import Path
import fitz  # PyMuPDF
import docx
from PIL import Image, ImageDraw, ImageFont

def generate_sample_data():
    base_dir = Path(__file__).resolve().parent
    academic_dir = base_dir / "authorized_academic"
    personal_dir = base_dir / "authorized_personal"
    images_dir = base_dir / "authorized_images"
    versions_dir = base_dir / "authorized_versions"
    unauthorized_dir = base_dir / "unauthorized_folder"

    for d in [academic_dir, personal_dir, images_dir, versions_dir, unauthorized_dir]:
        d.mkdir(parents=True, exist_ok=True)

    print("Generating Academic Documents...")
    # 1. Academic PDF: ML_Unit3_Notes.pdf
    pdf_path = academic_dir / "ML_Unit3_Notes.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 72),
        "Machine Learning Unit 3: Concept Learning & Candidate Elimination Algorithm\n\n"
        "1. Candidate Elimination Algorithm:\n"
        "The Candidate Elimination algorithm computes the version space containing all hypotheses\n"
        "from the hypothesis space H that are consistent with an observed sequence of training examples.\n"
        "It initializes the Specific boundary S to the most specific hypothesis and the General boundary G\n"
        "to the most general hypothesis.\n\n"
        "2. Version Space Representation:\n"
        "A hypothesis h is consistent with a set of training examples D if and only if h(x) = c(x)\n"
        "for each example (x, c(x)) in D. As new positive examples arrive, S is generalized;\n"
        "as negative examples arrive, G is specialized.\n\n"
        "3. Key Takeaway:\n"
        "Candidate elimination efficiently prunes hypotheses without enumerating all possible concepts.",
        fontsize=12
    )
    doc.save(str(pdf_path))
    doc.close()

    # 2. Academic DOCX: DBMS_Lecture_4.docx
    docx_path = academic_dir / "DBMS_Lecture_4.docx"
    doc = docx.Document()
    doc.add_heading("Database Management Systems: Relational Normalization", level=1)
    doc.add_paragraph(
        "Database normalization is the systematic approach of decomposing tables to eliminate data redundancy "
        "and undesirable characteristics like insertion, update, and deletion anomalies."
    )
    doc.add_heading("Normal Forms Summary:", level=2)
    p = doc.add_paragraph()
    p.add_run("• First Normal Form (1NF): Each column contains atomic values, no repeating groups.\n")
    p.add_run("• Second Normal Form (2NF): Must be in 1NF and have no partial functional dependency.\n")
    p.add_run("• Third Normal Form (3NF): Must be in 2NF and have no transitive dependency.\n")
    p.add_run("• Boyce-Codd Normal Form (BCNF): Strict version of 3NF where every determinant is a candidate key.")
    doc.save(str(docx_path))

    # 3. Academic TXT: Java_Concurrency_Notes.txt
    txt_path = academic_dir / "Java_Concurrency_Notes.txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(
            "Java Multi-Threading and Concurrency Notes\n"
            "==========================================\n\n"
            "Key topics:\n"
            "- Java Thread lifecycle: NEW, RUNNABLE, BLOCKED, WAITING, TIMED_WAITING, TERMINATED.\n"
            "- Synchronized methods and critical sections using intrinsic monitor locks.\n"
            "- Deadlock conditions: Mutual exclusion, Hold and wait, No preemption, Circular wait.\n"
            "- Executors framework and ThreadPoolExecutor for scalable thread management in backend services.\n"
            "- Exception handling in Java: Try, catch, finally, and custom exceptions.\n"
        )

    print("Generating Personal & Sensitive Documents (for Privacy testing)...")
    # 4. Personal PDF with PII (Email, Phone, Aadhaar): Student_Admission_Form.pdf
    personal_pdf = personal_dir / "Student_Admission_Form.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 72),
        "RAMDEOBABA UNIVERSITY - STUDENT REGISTRATION FORM\n"
        "--------------------------------------------------\n\n"
        "Applicant Name: Sample Student\n"
        "Department: Computer Science and Engineering\n"
        "Contact Email: student@example.com\n"
        "Mobile Number: +91 9876543210\n"
        "Government Identity (Aadhaar): 4521 8963 1254\n"
        "Permanent Address: Civil Lines, Nagpur, Maharashtra, India\n\n"
        "Declared that all educational credentials submitted for Semester V are authentic.",
        fontsize=12
    )
    doc.save(str(personal_pdf))
    doc.close()

    # 5. Personal DOCX with PII (PAN, Account): Bank_Statement_Summary.docx
    personal_docx = personal_dir / "Bank_Statement_Summary.docx"
    doc = docx.Document()
    doc.add_heading("Confidential Financial Statement Summary", level=1)
    doc.add_paragraph("Account Holder: Sample Account Holder")
    doc.add_paragraph("Tax Identity (PAN): ABCDE1234F")
    doc.add_paragraph("Email Address: holder@example.com")
    doc.add_paragraph("Primary Phone: 9812345678")
    doc.add_paragraph("Monthly scholarship stipend and university project allowance credits.")
    doc.save(str(personal_docx))

    # 6. Confidential Secrets File: Master_Password_Vault.txt
    pass_vault = personal_dir / "Master_Password_Vault.txt"
    with open(pass_vault, "w", encoding="utf-8") as f:
        f.write(
            "CONFIDENTIAL MASTER CREDENTIALS FILE\n"
            "====================================\n\n"
            "master_key: S3cur3P@ssw0rd!2026\n"
            "api_secret: sec_live_9948271049281726\n"
            "recovery_seed: yellow summer whisper mountain ocean galaxy\n"
            "WARNING: Do not distribute or index this document into search vectors.\n"
        )

    print("Generating Images for Multimodal & OCR Search...")
    # 7. Image: red car parked near a building: car_parked.jpg
    img_car = Image.new("RGB", (640, 480), color=(180, 210, 240))  # Sky/light background
    draw = ImageDraw.Draw(img_car)
    # Draw building
    draw.rectangle([50, 100, 300, 400], fill=(120, 120, 130), outline=(80, 80, 90))
    for r in range(120, 360, 50):
        for c in range(70, 270, 50):
            draw.rectangle([c, r, c + 30, r + 30], fill=(220, 240, 255))
    # Draw ground
    draw.rectangle([0, 400, 640, 480], fill=(70, 75, 80))
    # Draw Red Car
    draw.rectangle([340, 360, 580, 430], fill=(220, 30, 30), outline=(150, 0, 0))  # Car body
    draw.polygon([(380, 360), (430, 310), (520, 310), (550, 360)], fill=(200, 25, 25))  # Car cabin
    draw.polygon([(390, 355), (435, 315), (475, 315), (475, 355)], fill=(180, 220, 250))  # Window 1
    draw.polygon([(485, 315), (515, 315), (540, 355), (485, 355)], fill=(180, 220, 250))  # Window 2
    draw.ellipse([370, 410, 420, 460], fill=(20, 20, 20))  # Wheel 1
    draw.ellipse([500, 410, 550, 460], fill=(20, 20, 20))  # Wheel 2
    draw.text((360, 465), "Red car near building", fill=(255, 255, 255))
    img_car.save(str(images_dir / "car_parked.jpg"))

    # 8. Image: Classroom lecture: classroom_lecture.jpg
    img_class = Image.new("RGB", (640, 480), color=(245, 245, 240))
    draw = ImageDraw.Draw(img_class)
    # Whiteboard
    draw.rectangle([100, 50, 540, 220], fill=(255, 255, 255), outline=(100, 100, 100), width=4)
    draw.text((130, 80), "COMPUTER SCIENCE LECTURE HALL", fill=(40, 40, 150))
    draw.text((130, 120), "Today's Agenda: Multimodal Semantic Search & FAISS", fill=(30, 30, 30))
    draw.text((130, 160), "IDEA Lab Semester V - Ramdeobaba University", fill=(30, 30, 30))
    # Desks / Podiums
    draw.rectangle([50, 300, 250, 420], fill=(160, 110, 70))
    draw.rectangle([350, 300, 550, 420], fill=(160, 110, 70))
    img_class.save(str(images_dir / "classroom_lecture.jpg"))

    # 9. Screenshot with Java NullPointerException: java_error_screenshot.png
    img_err = Image.new("RGB", (700, 400), color=(30, 30, 30))  # Dark console
    draw = ImageDraw.Draw(img_err)
    draw.rectangle([0, 0, 700, 35], fill=(50, 50, 50))
    draw.text((15, 10), "Terminal - Java Application Execution", fill=(200, 200, 200))
    draw.text((20, 60), "$ java -jar build/libs/SeekVaultCore.jar", fill=(100, 220, 100))
    draw.text((20, 100), "[INFO] Starting SeekVault indexing daemon...", fill=(240, 240, 240))
    draw.text((20, 130), "[INFO] Connecting to SQLite local database: seekvault.db", fill=(240, 240, 240))
    draw.text((20, 170), "Exception in thread 'main' java.lang.NullPointerException", fill=(255, 80, 80))
    draw.text((40, 200), "at com.seekvault.indexer.DocumentScanner.parse(DocumentScanner.java:88)", fill=(255, 120, 120))
    draw.text((40, 230), "at com.seekvault.indexer.Pipeline.execute(Pipeline.java:34)", fill=(255, 120, 120))
    draw.text((40, 260), "at com.seekvault.Main.main(Main.java:15)", fill=(255, 120, 120))
    draw.text((20, 310), "Process terminated with exit code 1", fill=(200, 200, 200))
    img_err.save(str(images_dir / "java_error_screenshot.png"))

    print("Generating Versions & Duplicates...")
    # 10. Document Version 1: Project_Report_Draft_v1.docx
    doc_v1 = docx.Document()
    doc_v1.add_heading("SeekVault Project Report: Initial Draft", level=1)
    doc_v1.add_paragraph(
        "Abstract: SeekVault is a local-first semantic desktop file search tool. "
        "Traditional file search fails when filenames are generic. "
        "In this work, we propose vector embeddings for finding files."
    )
    doc_v1.add_paragraph("Authors: Abha Bhondge, Bhoomi Batra, Harsh Maske, Aarya Rohikar.")
    doc_v1.save(str(versions_dir / "Project_Report_Draft_v1.docx"))

    # 11. Document Version 2: Project_Report_Draft_v2.docx
    doc_v2 = docx.Document()
    doc_v2.add_heading("SeekVault Project Report: Revised Draft", level=1)
    doc_v2.add_paragraph(
        "Abstract: SeekVault is a local-first semantic desktop file search tool. "
        "Traditional file search fails when filenames are generic. "
        "In this work, we propose vector embeddings for finding files."
    )
    doc_v2.add_paragraph("Authors: Abha Bhondge, Bhoomi Batra, Harsh Maske, Aarya Rohikar.")
    doc_v2.add_heading("Detailed Methodology:", level=2)
    doc_v2.add_paragraph(
        "The system incorporates Microsoft Presidio for automated PII detection, "
        "tiered redaction for vector indices, and FAISS for sub-millisecond similarity queries."
    )
    doc_v2.save(str(versions_dir / "Project_Report_Draft_v2.docx"))

    # 12. Exact Duplicate: ML_Unit3_Notes_Copy.pdf (Exact duplicate of ML_Unit3_Notes.pdf)
    with open(academic_dir / "ML_Unit3_Notes.pdf", "rb") as src:
        content = src.read()
    with open(versions_dir / "ML_Unit3_Notes_Copy.pdf", "wb") as dst:
        dst.write(content)

    print("Generating Unauthorized Folder...")
    # 13. Unauthorized document: Secret_Research.pdf
    doc_unauth = fitz.open()
    p_un = doc_unauth.new_page()
    p_un.insert_text(
        (50, 72),
        "RESTRICTED RESEARCH DOCUMENT - OUTSIDE AUTHORIZATION\n"
        "---------------------------------------------------\n\n"
        "This file is placed inside sample_data/unauthorized_folder.\n"
        "SeekVault must strictly NEVER index or display this file in search results,\n"
        "proving that user folder authorization is enforced both during indexing\n"
        "and at query time.",
        fontsize=12
    )
    doc_unauth.save(str(unauthorized_dir / "Secret_Research.pdf"))
    doc_unauth.close()

    print("Sample data creation complete!")

if __name__ == "__main__":
    generate_sample_data()
