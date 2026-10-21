import os
from pathlib import Path
from typing import Dict, Any, Optional, List
import fitz  # PyMuPDF
import docx
import logging

logger = logging.getLogger(__name__)

class TextExtractor:
    @staticmethod
    def chunk_text(text: str, chunk_size: int = 250, overlap: int = 50) -> List[str]:
        """
        Split text into overlapping word chunks to ensure long documents
        fit within embedding model context limits without losing deep-page content.

        Args:
            text: Input document text.
            chunk_size: Target number of words per chunk (default 250 words).
            overlap: Overlapping words between consecutive chunks (default 50 words).

        Returns:
            List of chunk strings. If text has <= chunk_size words, returns [text].
        """
        if not text or not text.strip():
            return []

        words = text.split()
        if len(words) <= chunk_size:
            return [text.strip()]

        step = max(1, chunk_size - overlap)
        chunks = []
        for i in range(0, len(words), step):
            chunk = " ".join(words[i:i + chunk_size])
            if chunk.strip():
                chunks.append(chunk.strip())
            if i + chunk_size >= len(words):
                break

        return chunks
    @staticmethod
    def extract_pdf(file_path: Path) -> Dict[str, Any]:
        """Extract text and metadata from PDF files using PyMuPDF."""
        text_content = []
        page_count = 0
        try:
            doc = fitz.open(file_path)
            page_count = len(doc)
            for page_num in range(page_count):
                page = doc[page_num]
                page_text = page.get_text()
                if page_text:
                    text_content.append(page_text.strip())
            doc.close()
            full_text = "\n\n".join(text_content).strip()
            return {
                "success": True,
                "text": full_text,
                "page_count": page_count,
                "error": None
            }
        except Exception as e:
            logger.error(f"Error extracting PDF {file_path}: {e}")
            return {
                "success": False,
                "text": "",
                "page_count": page_count,
                "error": str(e)
            }

    @staticmethod
    def extract_docx(file_path: Path) -> Dict[str, Any]:
        """Extract text and tables from DOCX files using python-docx."""
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            
            # Also extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n\n".join(paragraphs).strip()
            return {
                "success": True,
                "text": full_text,
                "paragraph_count": len(paragraphs),
                "error": None
            }
        except Exception as e:
            logger.error(f"Error extracting DOCX {file_path}: {e}")
            return {
                "success": False,
                "text": "",
                "paragraph_count": 0,
                "error": str(e)
            }

    @staticmethod
    def extract_plain_text(file_path: Path) -> Dict[str, Any]:
        """Extract text from plain text files with UTF-8/fallback encoding."""
        encodings = ["utf-8", "utf-16", "latin-1", "cp1252"]
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    text = f.read().strip()
                return {
                    "success": True,
                    "text": text,
                    "encoding": enc,
                    "error": None
                }
            except (UnicodeDecodeError, Exception):
                continue
        
        return {
            "success": False,
            "text": "",
            "encoding": None,
            "error": "Failed to decode with known encodings"
        }

    @classmethod
    def extract(cls, file_path: Path) -> Dict[str, Any]:
        ext = file_path.suffix.lower()
        if ext == ".pdf":
            return cls.extract_pdf(file_path)
        elif ext == ".docx":
            return cls.extract_docx(file_path)
        elif ext in {".txt", ".md", ".log", ".csv", ".json", ".py", ".java"}:
            return cls.extract_plain_text(file_path)
        else:
            return {
                "success": False,
                "text": "",
                "error": f"Unsupported document extension: {ext}"
            }
